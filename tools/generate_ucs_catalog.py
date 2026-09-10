#!/usr/bin/env python3
"""Build PsyReaSFX's compact UCS catalog from the official translation workbook."""

from __future__ import annotations

import argparse
import hashlib
import os
from pathlib import Path

from openpyxl import load_workbook


FIELDS = (
    "category",
    "subcategory",
    "catid",
    "catshort",
    "explanation",
    "synonyms_en",
    "category_zh",
    "subcategory_zh",
    "synonyms_zh",
)

SOURCE_COLUMNS = (0, 1, 2, 3, 4, 5, 21, 22, 23)


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def escape_tsv(value: object) -> str:
    text = "" if value is None else str(value).strip()
    return (
        text.replace("%", "%25")
        .replace("\t", "%09")
        .replace("\r", "%0D")
        .replace("\n", "%0A")
    )


def build_catalog(source: Path, output: Path, expected_sha256: str) -> int:
    actual_sha256 = file_sha256(source)
    if expected_sha256 and actual_sha256 != expected_sha256.upper():
        raise ValueError(
            f"official workbook SHA-256 mismatch: {actual_sha256}"
        )

    workbook = load_workbook(source, read_only=True, data_only=True)
    worksheet = workbook.active
    headers = [cell.value for cell in next(
        worksheet.iter_rows(min_row=3, max_row=3)
    )]
    required_headers = {
        "Category",
        "SubCategory",
        "CatID",
        "CatShort",
        "Explanations",
        "Synonyms - Comma Separated",
        "Category_zh",
        "SubCategory_zh",
        "Synonyms_zh",
    }
    if not required_headers.issubset(set(headers)):
        missing = sorted(required_headers.difference(set(headers)))
        raise ValueError(f"official workbook columns missing: {missing}")

    records: list[tuple[str, ...]] = []
    seen_catids: set[str] = set()
    for source_row in worksheet.iter_rows(min_row=4, values_only=True):
        values = tuple(
            "" if source_row[index] is None else str(source_row[index]).strip()
            for index in SOURCE_COLUMNS
        )
        catid = values[2]
        if not catid:
            continue
        if catid in seen_catids:
            raise ValueError(f"duplicate CatID in official workbook: {catid}")
        if not values[0] or not values[1]:
            raise ValueError(f"incomplete UCS row: {catid}")
        seen_catids.add(catid)
        records.append(values)

    if len(records) != 753:
        raise ValueError(f"expected 753 UCS 8.2.1 rows, found {len(records)}")

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(output.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8", newline="\n") as handle:
        handle.write("#schema\tucs_catalog_v1\n")
        handle.write("#ucs_version\t8.2.1\n")
        handle.write("#source_file\tUCS v8.2.1 Full Translations.xlsx\n")
        handle.write(f"#source_sha256\t{actual_sha256}\n")
        handle.write("\t".join(FIELDS) + "\n")
        for record in records:
            handle.write("\t".join(escape_tsv(value) for value in record) + "\n")
    os.replace(temporary, output)
    return len(records)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--expected-sha256", default="")
    args = parser.parse_args()
    count = build_catalog(args.source, args.output, args.expected_sha256)
    print(f"Generated {args.output}: {count} UCS records")


if __name__ == "__main__":
    main()
